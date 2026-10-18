import asyncio
import io
import json
import os
import tempfile
import unittest
import uuid
from contextlib import contextmanager
from pathlib import Path
from unittest.mock import patch

import httpx
from docx import Document
from langchain_core.messages import AIMessage, HumanMessage, ToolMessage
from langchain_core.language_models.fake_chat_models import FakeMessagesListChatModel
from langgraph.prebuilt import create_react_agent
from pypdf import PdfWriter
from pypdf.generic import DecodedStreamObject, DictionaryObject, NameObject

from agent.db import psycopg, sql, _build_dsn
from agent.knowledge import (
    EmbeddingClient, EmbeddingSettings, KnowledgeBase, KnowledgeError, NeedsOCR, RAG_INSTRUCTIONS,
    chunk_sections, extract_rule_references, make_search_tool, parse_document,
)


def word_bytes(text="Flood response: protect hospitals and evacuate flooded roads."):
    doc = Document()
    doc.add_heading("Emergency response rules", 0)
    doc.add_paragraph(text)
    table = doc.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Water depth (m)"
    table.cell(0, 1).text = "Required action"
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


class ParserTests(unittest.TestCase):
    def test_docx_paragraphs_and_tables_keep_units(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "rule.docx"
            path.write_bytes(word_bytes())
            sections = parse_document(path)
            self.assertIn("Water depth (m)", "\n".join(text for _, text in sections))
            self.assertTrue(all("正文" in locator for locator, _ in sections))

    def test_pdf_preserves_real_page_numbers(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "rule.pdf"
            writer = PdfWriter()
            for text in ["First page flood rules", "Second page emergency routes"]:
                page = writer.add_blank_page(width=200, height=100)
                font = DictionaryObject({NameObject('/Type'): NameObject('/Font'),
                                         NameObject('/Subtype'): NameObject('/Type1'),
                                         NameObject('/BaseFont'): NameObject('/Helvetica')})
                page[NameObject('/Resources')] = DictionaryObject({
                    NameObject('/Font'): DictionaryObject({NameObject('/F1'): font}),
                })
                stream = DecodedStreamObject()
                stream.set_data(f'BT /F1 12 Tf 10 50 Td ({text}) Tj ET'.encode())
                page[NameObject('/Contents')] = stream
            writer.write(path)
            self.assertEqual([loc for loc, _ in parse_document(path)], ["第 1 页", "第 2 页"])

    def test_image_only_pdf_requires_ocr(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "scan.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            writer.write(path)
            with self.assertRaises(NeedsOCR):
                parse_document(path)

    def test_chunks_are_bounded_and_keep_source_locator(self):
        text = "洪水处置规则：道路受淹后先核实水深，不能只根据面积判断危险。" * 100
        chunks = chunk_sections([("第 7 页", text)])
        self.assertGreater(len(chunks), 1)
        self.assertTrue(all(len(c["content"]) <= 900 and c["locator"] == "第 7 页" for c in chunks))


class EmbeddingTests(unittest.TestCase):
    def test_independent_endpoint_never_receives_chat_key(self):
        with patch.dict(os.environ, {"RAG_EMBEDDING_BASE_URL": "https://other.example/v1"}, clear=True):
            settings = EmbeddingSettings.from_env({
                "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1", "api_key": "chat-secret",
            })
        self.assertEqual(settings.api_key, "")
        self.assertFalse(settings.configured)

    def test_native_endpoint_can_reuse_existing_configuration(self):
        with patch.dict(os.environ, {}, clear=True):
            settings = EmbeddingSettings.from_env({
                "base_url": "https://workspace.cn-beijing.maas.aliyuncs.com/compatible-mode/v1", "api_key": "chat-secret",
            })
        self.assertTrue(settings.configured)
        self.assertEqual(settings.dimensions, 1024)
        self.assertEqual(settings.model, "qwen3.7-text-embedding-flash")

    def test_explicit_embedding_model_remains_supported(self):
        with patch.dict(os.environ, {"RAG_EMBEDDING_MODEL": "text-embedding-v4"}, clear=True):
            settings = EmbeddingSettings.from_env({
                "base_url": "https://dashscope.aliyuncs.com/compatible-mode/v1", "api_key": "chat-secret",
            })
        self.assertEqual(settings.model, "text-embedding-v4")
        flash = EmbeddingSettings(settings.base_url, settings.api_key)
        self.assertNotEqual(settings.profile, flash.profile)

    def test_flash_openai_compatible_payload_and_vector_validation(self):
        client = EmbeddingClient(EmbeddingSettings("https://example.test/v1", "secret"))
        seen = []

        def handle(request):
            seen.append(json.loads(request.content))
            return httpx.Response(200, json={"data": [
                {"index": 0, "embedding": [3, 4] + [0] * 1022},
            ]})

        original_client = httpx.Client
        with patch("agent.knowledge.httpx.Client", side_effect=lambda **kw: original_client(
            transport=httpx.MockTransport(handle), **kw,
        )):
            vector = client.embed(["test document"])[0]
            query = client.embed(["test query"], query=True)[0]
        self.assertEqual(len(vector), 1024)
        self.assertEqual(vector[:2], [.6, .8])
        self.assertEqual(vector, query)
        self.assertTrue(all(body["model"] == "qwen3.7-text-embedding-flash" for body in seen))
        self.assertTrue(all(body["dimensions"] == 1024 for body in seen))
        self.assertTrue(all(body["encoding_format"] == "float" for body in seen))
        self.assertTrue(all("text_type" not in body for body in seen))

    def test_embedding_order_normalization_and_query_type(self):
        client = EmbeddingClient(EmbeddingSettings("https://example.test/v1", "secret", model="text-embedding-v4", dimensions=2))
        seen = []

        def handle(request):
            seen.append(json.loads(request.content))
            return httpx.Response(200, json={"data": [
                {"index": 1, "embedding": [0, 5]}, {"index": 0, "embedding": [3, 4]},
            ]})

        original_client = httpx.Client
        with patch("agent.knowledge.httpx.Client", side_effect=lambda **kw: original_client(transport=httpx.MockTransport(handle), **kw)):
            vectors = client.embed(["a", "b"], query=True)
        self.assertEqual(vectors, [[.6, .8], [0, 1]])
        self.assertEqual(seen[0]["text_type"], "query")

    def test_quota_failure_is_not_reported_as_success(self):
        client = EmbeddingClient(EmbeddingSettings("https://example.test/v1", "secret", dimensions=2))
        original_client = httpx.Client
        with patch("agent.knowledge.httpx.Client", side_effect=lambda **kw: original_client(
            transport=httpx.MockTransport(lambda req: httpx.Response(403, json={"error": {"code": "insufficient_quota"}})), **kw,
        )):
            with self.assertRaisesRegex(KnowledgeError, "额度不足"):
                client.embed(["test"])

    def test_invalid_vector_rejected(self):
        client = EmbeddingClient(EmbeddingSettings("https://example.test/v1", "secret", dimensions=2))
        original_client = httpx.Client
        with patch("agent.knowledge.httpx.Client", side_effect=lambda **kw: original_client(
            transport=httpx.MockTransport(lambda req: httpx.Response(200, json={"data": [{"index": 0, "embedding": [0, 0]}]})), **kw,
        )):
            with self.assertRaisesRegex(KnowledgeError, "无效"):
                client.embed(["test"])


class RuleToolTests(unittest.TestCase):
    def test_report_questions_are_in_scope_and_knowledge_precedes_web_search(self):
        self.assertIn("disaster reports", RAG_INSTRUCTIONS)
        self.assertIn("before public web search", RAG_INSTRUCTIONS)
        self.assertIn("not restricted to domestic", RAG_INSTRUCTIONS)

    def test_followup_history_is_hidden_from_model_and_forwarded_to_search(self):
        class Probe:
            def search(self, *args, **kwargs):
                return {"query": args[0], **kwargs}
        tool = make_search_tool(Probe())
        self.assertNotIn("config", tool.args)
        self.assertNotIn("history", tool.args)
        result = json.loads(tool.invoke({"query": "委内瑞拉地震情况"}, config={
            "configurable": {
                "rag_question": "RAG库里不是有吗",
                "rag_history": ["你知道委内瑞拉地震情况吗"],
            },
        }))
        self.assertEqual(result["question"], "RAG库里不是有吗")
        self.assertEqual(result["history"], ["你知道委内瑞拉地震情况吗"])

    def test_original_question_is_injected_without_exposing_config_to_model(self):
        class Probe:
            def search(self, *args, **kwargs):
                return {"results": [], "original": kwargs["question"], "query": args[0]}
        tool = make_search_tool(Probe())
        self.assertNotIn("config", tool.args)
        result = json.loads(tool.invoke({"query": "compressed keywords"}, config={
            "configurable": {"rag_question": "What does my test document actually say?"},
        }))
        self.assertEqual(result["query"], "compressed keywords")
        self.assertEqual(result["original"], "What does my test document actually say?")

    def test_only_cited_retrieved_evidence_becomes_reference(self):
        refs = [{"citation": "[规则-a]", "id": "a"}, {"citation": "[规则-b]", "id": "b"}]
        response = {"messages": [ToolMessage(
            content=json.dumps({"results": refs, "status": "sufficient"}), name="search_rule_documents", tool_call_id="search",
        )]}
        self.assertEqual(extract_rule_references(response, "依据[规则-a]。"), [refs[0]])

    def test_rejected_evidence_cannot_be_saved_as_a_reference(self):
        response = {"messages": [ToolMessage(
            content=json.dumps({"status": "insufficient", "results": [{"citation": "[规则-a]"}]}),
            name="search_rule_documents", tool_call_id="search",
        )]}
        self.assertEqual(extract_rule_references(response, "依据[规则-a]。"), [])

    def test_retrieval_failure_is_explicit_to_agent(self):
        class Missing:
            def search(self, *args, **kwargs):
                raise KnowledgeError("Embedding 额度不足")
        payload = json.loads(make_search_tool(Missing()).invoke({"query": "洪水规则"}))
        self.assertEqual(payload["results"], [])
        self.assertIn("额度不足", payload["error"])


class BackendCitationTests(unittest.TestCase):
    def test_retrieval_config_only_keeps_recent_user_text_without_attachments(self):
        import backend_api
        history = [
            {"role": "user", "content": "old question " + str(i)} for i in range(6)
        ] + [
            {"role": "user", "content": "你知道委内瑞拉地震情况吗"},
            {"role": "assistant", "content": "There are no overseas reports. Incorrect prior answer."},
            {"role": "user", "content": [
                {"type": "text", "text": "分析这份报告\n\nUploaded files:\n/private/report.docx"},
                {"type": "image_url", "image_url": {"url": "data:image/png;base64,secret"}},
            ]},
            {"role": "user", "content": "RAG库里不是有吗"},
        ]
        config = backend_api.build_retrieval_config("RAG库里不是有吗", history)
        self.assertEqual(config["rag_question"], "RAG库里不是有吗")
        self.assertEqual(len(config["rag_history"]), 4)
        self.assertIn("你知道委内瑞拉地震情况吗", config["rag_history"])
        self.assertIn("分析这份报告", config["rag_history"])
        self.assertNotIn("RAG库里不是有吗", config["rag_history"])
        self.assertNotIn("Incorrect", json.dumps(config))
        self.assertNotIn("secret", json.dumps(config))
        self.assertNotIn("/private/", json.dumps(config))
        self.assertTrue(all(len(text) <= 800 for text in config["rag_history"]))

    def test_session_restores_database_text_once_after_backend_restart(self):
        import backend_api
        rows = [
            {"role": "user", "content": "你知道委内瑞拉地震情况吗", "display_content": "你知道委内瑞拉地震情况吗"},
            {"role": "assistant", "content": "Previous answer", "display_content": "Previous answer"},
        ]
        with tempfile.TemporaryDirectory() as root:
            with (
                patch.object(backend_api, "TEMP_BASE", Path(root)),
                patch.object(backend_api, "SESSIONS", {}),
                patch.object(backend_api.db, "get_chat_context", return_value=rows) as load,
            ):
                session = backend_api.get_session("restored-session")
                self.assertEqual(len(session.messages), 2)
                self.assertIs(backend_api.get_session("restored-session"), session)
                load.assert_called_once_with("restored-session")
                session.messages.append({"role": "user", "content": "RAG库里不是有吗"})
                config = backend_api.build_retrieval_config("RAG库里不是有吗", session.messages)
                self.assertEqual(config["rag_history"], ["你知道委内瑞拉地震情况吗"])

    def test_normal_and_stream_chat_return_and_save_rule_references(self):
        import backend_api
        reference = {"id": "chunk-test", "document_id": str(uuid.uuid4()),
                     "filename": "防汛规则.docx", "citation": "[规则-test]", "locator": "第 1 页",
                     "content": "保障医院通行。", "url": "/api/knowledge/documents/test/file"}
        response = {"messages": [
            ToolMessage(content=json.dumps({"results": [reference], "status": "sufficient"}), name="search_rule_documents", tool_call_id="rules-test"),
            AIMessage(content="<Conclusion>应保障医院通行[规则-test]。</Conclusion>"),
        ]}

        class Handle:
            def invoke(self, *args, **kwargs):
                assert kwargs["config"]["configurable"]["rag_question"] == "防汛规则"
                assert kwargs["config"]["configurable"]["rag_history"] == ["此前询问洪水处置要求"]
                return response

            def stream(self, *args, **kwargs):
                assert kwargs["config"]["configurable"]["rag_question"] == "防汛规则"
                assert kwargs["config"]["configurable"]["rag_history"] == ["此前询问洪水处置要求"]
                yield {"type": "final", "response": response}

            def review_answer_and_artifacts(self, raw, uploaded, outputs, **kwargs):
                assert kwargs["rule_references"] == [reference]
                return raw

        for streaming in (False, True):
            with self.subTest(streaming=streaming), tempfile.TemporaryDirectory() as root:
                session = backend_api.ChatSession(session_id="test", temp_dir=Path(root), handle=Handle())
                session.messages.append({"role": "user", "content": "此前询问洪水处置要求"})
                with (
                    patch.object(backend_api, "get_session", return_value=session),
                    patch.object(backend_api.db, "create_session"),
                    patch.object(backend_api.db, "save_chat_message") as save,
                    patch.object(backend_api, "maybe_generate_session_title"),
                ):
                    fn = backend_api.chat_stream if streaming else backend_api.chat
                    result = fn(session_id="test", message="防汛规则", system_prompt="Test",
                                recursion_limit=40, max_execution_time=600, show_trace=False, files=None)
                    if streaming:
                        async def consume():
                            return [chunk async for chunk in result.body_iterator]
                        done = next(chunk for chunk in asyncio.run(consume()) if "event: done" in chunk)
                        payload = json.loads(done.split("data: ", 1)[1])
                    else:
                        payload = result
                    self.assertEqual(payload["retrieval_refs"], [reference])
                    self.assertEqual(save.call_args.kwargs["retrieval_refs"], [reference])
                    self.assertEqual(backend_api._serialize_message({
                        "id": 1, "role": "assistant", "content": payload["answer"], "retrieval_refs": [reference],
                    })["retrieval_refs"], [reference])


class FakeEmbedder:
    """Deterministic vectors for integration tests, never used in production."""
    def __init__(self):
        self.calls = 0

    def embed(self, texts, *, query=False):
        self.calls += 1
        return [[1., 0.] if "flood" in text.lower() else [0., 1.] for text in texts]


class FakeVerifier:
    """Isolate storage tests from both paid verification services."""
    def verify(self, query, candidates, top_k, region, disaster_type):
        return {"results": candidates[:top_k], "status": "sufficient", "message": "Test evidence accepted"}

    def status(self):
        return {"rerank_model": "fake", "judge_model": "fake", "rerank_configured": True, "judge_configured": True}


@unittest.skipUnless(os.getenv("RAG_TEST_DB") == "1", "Set RAG_TEST_DB=1 for isolated PostgreSQL tests")
class DatabaseIntegrationTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        from dotenv import load_dotenv
        load_dotenv(".env", override=True)
        cls.schema = "test_knowledge_" + uuid.uuid4().hex
        cls.dsn = _build_dsn()
        with psycopg.connect(cls.dsn, connect_timeout=3) as conn:
            conn.execute(sql.SQL("CREATE SCHEMA {}").format(sql.Identifier(cls.schema)))

    @classmethod
    def tearDownClass(cls):
        with psycopg.connect(cls.dsn, connect_timeout=3) as conn:
            conn.execute(sql.SQL("DROP SCHEMA {} CASCADE").format(sql.Identifier(cls.schema)))

    def setUp(self):
        test = self

        class TestDatabase:
            def health_check(self):
                return True

            @contextmanager
            def _conn(self):
                with psycopg.connect(test.dsn, connect_timeout=3) as conn:
                    conn.execute(sql.SQL("SET search_path TO {}").format(sql.Identifier(test.schema)))
                    yield conn

        self.database = TestDatabase()
        with self.database._conn() as conn:
            conn.execute("CREATE TABLE IF NOT EXISTS chat_messages (id integer)")
        self.temporary = tempfile.TemporaryDirectory()
        self.addCleanup(self.temporary.cleanup)
        self.embedder = FakeEmbedder()
        self.knowledge = KnowledgeBase(Path(self.temporary.name), EmbeddingSettings(
            "https://test.invalid/v1", "test", dimensions=2,
        ), database=self.database, embedder=self.embedder, verifier=FakeVerifier())
        self.knowledge.ensure_schema()
        with self.database._conn() as conn:
            conn.execute("TRUNCATE knowledge_documents CASCADE")

    def add_document(self, filename="flood.docx", text="Flood response: protect hospitals.", **meta):
        result = self.knowledge.upload(filename, word_bytes(text), meta)
        self.assertTrue(self.knowledge.process_one())
        return result["document"]["id"]

    def test_ingest_duplicate_search_disable_delete(self):
        document_id = self.add_document(region="深圳", disaster_type="洪水")
        self.assertEqual(self.knowledge.get_document(document_id)["status"], "ready")
        duplicate = self.knowledge.upload("same.docx", self.knowledge.document_path(
            self.knowledge.get_document(document_id)).read_bytes(), {})
        self.assertTrue(duplicate["duplicate"])
        self.assertEqual(duplicate["document"]["id"], document_id)
        results = self.knowledge.search("flood emergency", region="深圳")["results"]
        self.assertTrue(results)
        self.assertEqual(results[0]["filename"], "flood.docx")
        self.assertIn("规则-", results[0]["citation"])
        self.assertEqual(self.knowledge.search("flood", region="北京")["results"], [])
        self.knowledge.update_document(document_id, enabled=False)
        self.assertEqual(self.knowledge.search("flood")["results"], [])
        self.knowledge.delete(document_id)
        self.assertFalse(list(Path(self.temporary.name).iterdir()))

    def test_no_library_and_no_match_do_not_force_results(self):
        self.assertEqual(self.knowledge.search("flood")["results"], [])
        self.assertEqual(self.embedder.calls, 0)
        self.add_document("earthquake.docx", "Earthquake shelter plan")
        self.assertEqual(self.knowledge.search("flood")["results"], [])

    def test_verification_rejection_and_outage_never_return_raw_candidates(self):
        from agent.rag_verification import RetrievalVerificationError
        self.add_document()
        with patch.object(self.knowledge.verifier, "verify", return_value={
            "results": [], "status": "insufficient", "message": "Incomplete evidence",
        }):
            result = self.knowledge.search("flood")
            self.assertEqual(result["results"], [])
            self.assertEqual(result["status"], "insufficient")
        with patch.object(self.knowledge.verifier, "verify", side_effect=RetrievalVerificationError("核验超时")):
            result = self.knowledge.search("flood")
            self.assertEqual(result["results"], [])
            self.assertEqual(result["status"], "unavailable")
            self.assertIn("核验超时", result["message"])

    def test_recall_pool_is_larger_than_final_result_limit(self):
        self.add_document("flood-1.docx", "Flood response: verify image dates.")
        self.add_document("flood-2.docx", "Flood response: verify coordinate reference systems.")
        with patch.object(self.knowledge.verifier, "verify", wraps=self.knowledge.verifier.verify) as verify:
            result = self.knowledge.search("flood", top_k=1)
        self.assertGreaterEqual(len(verify.call_args.args[1]), 2)
        self.assertEqual(len(result["results"]), 1)

    def test_full_question_not_keywords_controls_evidence_verification(self):
        self.add_document()
        original = "请读取测试文档中的 flood 报告核对要求，不是判断正式法规效力。"
        with patch.object(self.knowledge.verifier, "verify", wraps=self.knowledge.verifier.verify) as verify:
            self.knowledge.search("flood checklist", question=original)
        self.assertEqual(verify.call_args.args[0], original)

    def test_followup_context_reaches_reranking_and_evidence_verification(self):
        self.add_document()
        with patch.object(self.knowledge.verifier, "verify", wraps=self.knowledge.verifier.verify) as verify:
            self.knowledge.search("flood report", question="RAG库里不是有吗",
                                  history=["你知道洪水报告的内容吗"])
        context = json.loads(verify.call_args.args[0])
        self.assertEqual(context["current_question"], "RAG库里不是有吗")
        self.assertEqual(context["retrieval_intent"], "flood report")
        self.assertEqual(context["previous_user_questions"], ["你知道洪水报告的内容吗"])

    def test_context_is_bounded_and_current_question_keeps_its_constraints(self):
        self.add_document()
        original = "查询深圳洪水报告，但不要给出没有依据的水深。" + "请保留所有约束。" * 400
        history = [str(i) * 2000 for i in range(10)]
        with patch.object(self.knowledge.verifier, "verify", wraps=self.knowledge.verifier.verify) as verify:
            self.knowledge.search("flood report", question=original, history=history)
        context = json.loads(verify.call_args.args[0])
        self.assertEqual(context["current_question"], original)
        self.assertEqual(len(context["previous_user_questions"]), 4)
        self.assertTrue(all(len(text) <= 800 for text in context["previous_user_questions"]))

    def test_database_context_only_loads_recent_text_in_chronological_order(self):
        from agent.db import Database
        with self.database._conn() as conn:
            conn.execute("DROP TABLE chat_messages")
            conn.execute("""CREATE TABLE chat_messages (
                id SERIAL PRIMARY KEY, session_id TEXT, role TEXT, content TEXT,
                display_content TEXT, created_at TIMESTAMPTZ DEFAULT now(),
                attachment_files JSONB)""")
            for i in range(45):
                conn.execute("""INSERT INTO chat_messages
                    (session_id,role,content,display_content,attachment_files)
                    VALUES ('history-test','user',%s,%s,'[{"data_base64":"not-context"}]')""",
                             (f"question {i}", f"question {i}"))
        with patch.object(Database, "_conn", self.database._conn), patch.object(Database, "_ensure_pool", return_value=True):
            rows = Database().get_chat_context("history-test")
        self.assertEqual(len(rows), 40)
        self.assertEqual(rows[0]["content"], "question 5")
        self.assertEqual(rows[-1]["content"], "question 44")
        self.assertTrue(all("attachment_files" not in row for row in rows))

    def test_failure_atomicity_and_retry_after_restart(self):
        document = self.knowledge.upload("flood.docx", word_bytes(), {})["document"]
        with patch.object(self.embedder, "embed", side_effect=KnowledgeError("额度不足")):
            self.knowledge.process_one()
        self.assertEqual(self.knowledge.get_document(document["id"])["status"], "failed")
        self.assertEqual(self.knowledge.search("flood")["results"], [])
        self.knowledge.retry(document["id"])
        restarted = KnowledgeBase(Path(self.temporary.name), self.knowledge.settings,
                                  database=self.database, embedder=self.embedder)
        restarted.process_one()
        self.assertEqual(restarted.get_document(document["id"])["status"], "ready")

    def test_model_profile_change_requires_reindex(self):
        document_id = self.add_document()
        changed = KnowledgeBase(Path(self.temporary.name), EmbeddingSettings(
            "https://test.invalid/v1", "test", model="other-model", dimensions=2,
        ), database=self.database, embedder=self.embedder)
        self.assertTrue(changed.get_document(document_id)["needs_reindex"])
        calls_before = self.embedder.calls
        result = changed.search("flood")
        self.assertEqual(result["results"], [])
        self.assertEqual(result["status"], "unavailable")
        self.assertEqual(result["reason"], "index_needs_rebuild")
        self.assertEqual(result["documents_need_reindex"], 1)
        self.assertIn("重建索引", result["message"])
        self.assertEqual(self.embedder.calls, calls_before)

    def test_reindex_warning_respects_enabled_state_and_scope(self):
        document_id = self.add_document(region="Shenzhen")
        changed = KnowledgeBase(Path(self.temporary.name), EmbeddingSettings(
            "https://test.invalid/v1", "test", model="other-model", dimensions=2,
        ), database=self.database, embedder=self.embedder)
        self.assertEqual(changed.search("flood", region="Beijing")["status"], "no_match")
        changed.update_document(document_id, enabled=False)
        self.assertEqual(changed.search("flood")["status"], "no_match")

    def test_partial_reindex_does_not_report_missing_knowledge(self):
        self.add_document("old-flood.docx", "Flood response: verify image dates.")
        changed = KnowledgeBase(Path(self.temporary.name), EmbeddingSettings(
            "https://test.invalid/v1", "test", model="other-model", dimensions=2,
        ), database=self.database, embedder=self.embedder, verifier=FakeVerifier())
        changed.upload("new-earthquake.docx", word_bytes("Earthquake background."), {})
        changed.process_one()
        self.assertEqual(changed.search("flood")["reason"], "index_needs_rebuild")
        self.assertEqual(changed.search("earthquake")["status"], "sufficient")
        with patch.object(changed.verifier, "verify", return_value={
            "status": "insufficient", "results": [], "message": "No answer in indexed documents",
        }):
            self.assertEqual(changed.search("earthquake")["reason"], "index_needs_rebuild")

    def test_expired_lease_is_recovered_and_deleted_job_is_not_resurrected(self):
        document = self.knowledge.upload("flood.docx", word_bytes(), {})["document"]
        with self.database._conn() as conn:
            conn.execute("UPDATE knowledge_documents SET status='embedding',lease_until=now()-interval '1 minute' WHERE id=%s", (document["id"],))
        self.knowledge.process_one()
        self.assertEqual(self.knowledge.get_document(document["id"])["status"], "ready")
        self.knowledge.retry(document["id"])
        def delete_while_embedding(texts, **kwargs):
            self.knowledge.delete(document["id"])
            return [[1., 0.] for _ in texts]
        with patch.object(self.embedder, "embed", side_effect=delete_while_embedding):
            self.knowledge.process_one()
        self.assertEqual(self.knowledge.list_documents(), [])

    def test_scan_and_invalid_upload(self):
        with self.assertRaises(KnowledgeError):
            self.knowledge.upload("fake.pdf", b"not a pdf", {})
        with self.assertRaises(KnowledgeError):
            self.knowledge.upload("old.doc", b"old word", {})
        writer = PdfWriter()
        writer.add_blank_page(width=100, height=100)
        output = io.BytesIO()
        writer.write(output)
        document = self.knowledge.upload("scan.pdf", output.getvalue(), {})["document"]
        self.knowledge.process_one()
        self.assertEqual(self.knowledge.get_document(document["id"])["status"], "needs_ocr")

    def test_fastapi_crud_preview_and_search(self):
        from fastapi import FastAPI
        from fastapi.testclient import TestClient
        from agent.knowledge_api import knowledge_router
        app = FastAPI()
        app.include_router(knowledge_router(self.knowledge))
        client = TestClient(app)
        response = client.post('/api/knowledge/documents', files={"file": ("flood.docx", word_bytes())})
        self.assertEqual(response.status_code, 202)
        document_id = response.json()["document"]["id"]
        self.knowledge.process_one()
        self.assertTrue(client.get(f'/api/knowledge/documents/{document_id}').json()["preview"])
        self.assertTrue(client.post('/api/knowledge/search', json={"query": "flood"}).json()["results"])
        self.assertEqual(client.get(f'/api/knowledge/documents/{document_id}/file').status_code, 200)
        self.assertEqual(client.patch(f'/api/knowledge/documents/{document_id}', json={"enabled": False}).status_code, 200)
        self.assertEqual(client.delete(f'/api/knowledge/documents/{document_id}').status_code, 200)
        self.assertEqual(client.get(f'/api/knowledge/documents/{document_id}').status_code, 404)
        self.assertEqual(client.get('/api/knowledge/documents/not-uuid').status_code, 422)

    def test_real_langgraph_calls_retrieval_and_preserves_citations(self):
        document_id = self.add_document()
        reference = self.knowledge.search("flood")["results"][0]
        llm = FakeMessagesListChatModel(responses=[
            AIMessage(content="", tool_calls=[{"name": "search_rule_documents", "args": {"query": "flood"}, "id": "rules-call", "type": "tool_call"}]),
            AIMessage(content=f"<Conclusion>Protect hospitals {reference['citation']}</Conclusion>"),
        ])
        with patch.object(FakeMessagesListChatModel, "bind_tools", lambda model, tools, **kwargs: model, create=True):
            agent = create_react_agent(llm, [make_search_tool(self.knowledge)])
            with patch.object(self.knowledge.verifier, "verify", wraps=self.knowledge.verifier.verify) as verify:
                response = asyncio.run(agent.ainvoke({"messages": [HumanMessage(content="flood response rules?")]}, config={
                    "configurable": {"rag_question": "What does the uploaded document say about flood response?"},
                }))
                self.assertEqual(verify.call_args.args[0], "What does the uploaded document say about flood response?")
        refs = extract_rule_references(response, response["messages"][-1].content)
        self.assertEqual(refs[0]["document_id"], document_id)


if __name__ == '__main__':
    unittest.main()
